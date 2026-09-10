from pathlib import Path
import json,csv,hashlib,shutil,re,zipfile
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
ROOT=Path('/home/user');DEST=ROOT/'MOKO';ORG=DEST/'99_Технические_материалы/Организация'
ORG.mkdir(parents=True,exist_ok=True)
PLAN=ORG/'Карта_переноса.json'
YT='01_YouTube';TG='02_Telegram';GH='03_GitHub';VK='04_VK';BR='05_Бренд_MOKO';RAW='06_Материалы_от_вас';TECH='99_Технические_материалы/Рабочий_код_и_данные'
YT94=YT+'/01_Обложки_94';TGF=TG+'/01_Аватары_ФИНАЛ_В';EM=TG+'/02_Эмодзи_24';GHA=GH+'/01_Аватар_организации'

def route(old):
 p=Path(old);r=p.parts;ext=p.suffix.lower()
 if r[0]=='uploads':return RAW+'/'+str(Path(*r[1:]))
 rel=Path(*r[1:]);a=rel.parts
 # Data and generation sources keep their original hierarchy in the technical area.
 if ext in ['.py','.json','.log'] or 'source' in a or 'assets' in a or 'research' in a or 'briefs' in a:
  return TECH+'/'+old
 if a[0]=='brand':
  return BR+('/03_Шрифты/' if ext=='.ttf' else '/01_Исходный_знак/')+p.name
 if a[0]=='originals':return YT+'/80_Исходные_обложки_YouTube/'+p.name
 if a[:2]==('production','thumbnails'):return YT94+'/thumbnails/'+p.name
 if a[0]=='final':return YT94+'/'+str(Path(*a[1:]))
 if a[0]=='branding':
  if len(a)>1 and a[1]=='logo':return BR+'/02_Логотипы/'+p.name
  if len(a)>1 and a[1]=='avatars':return BR+'/04_Общие_аватары_YouTube_VK/'+p.name
  if len(a)>1 and a[1]=='youtube':return YT+'/02_Оформление_канала/'+p.name
  if len(a)>1 and a[1]=='guides':return YT+'/02_Оформление_канала/Проверка_обрезки_НЕ_загружать/'+p.name
  if len(a)>1 and a[1]=='streams':return YT+'/03_Заставки_для_эфиров/'+p.name
  if len(a)>1 and a[1]=='vk':
   return VK+('/02_Меню_8/' if len(a)>2 and a[2]=='menu' else '/01_Обложки/')+p.name
  if len(a)>1 and a[1]=='telegram':return TG+'/03_Баннер_закрепа_и_обои/'+p.name
  return BR+'/05_Базовый_комплект_оформления/'+p.name
 if a[0]=='emoji':return EM+'/'+str(Path(*a[1:]))
 if a[0]=='github':return GHA+'/'+str(Path(*a[1:]))
 if a[0]=='telegram_avatars':
  tail=list(a[1:]);first=tail[0]
  if first=='final_graphic':return TGF+'/'+str(Path(*tail[1:]))
  names={'concept02':'02_Крупные_символы','concept03':'03_Фирменная_текстура','alternatives04':'04_Выбор_символов','group05':'05_Варианты_группы','selected':'06_Промежуточный_комплект','bold_concepts':'07_Направления_А_Б_В','bc_variants':'08_Доработки_Б_и_В'}
  if first in names:return TG+'/90_Архив_аватаров/'+names[first]+'/'+str(Path(*tail[1:]))
  return TG+'/90_Архив_аватаров/01_Первые_аватары/'+str(Path(*tail))
 if a[0]=='docs':return YT+'/90_Архив_этапов/12_Старые_документы/'+p.name
 if a[0]=='releases':return YT+'/90_Архив_этапов/11_Выпуск_18/'+p.name
 if a[0]=='production':return YT+'/90_Архив_этапов/10_Выпуск_12/'+p.name
 if a[0]=='revision10':return YT+'/81_Согласованные_портреты/Стрим_v10/'+p.name
 if a[0]=='revision12':return YT+'/81_Согласованные_портреты/Expo_v12/'+p.name
 names={'concepts':'01_Первые_концепты','selected':'02_Первые_6_обложек','branded':'03_Пример_брендирования','revision2':'04_Серия_v2','revision3':'05_Серия_v3','batch02':'08_Партия_02','batch03':'09_Партия_03'}
 if a[0] in names:return YT+'/90_Архив_этапов/'+names[a[0]]+'/'+str(Path(*a[1:]))
 if a[0].startswith('revision'):return YT+'/90_Архив_этапов/06_Поиск_образа/'+str(rel)
 if ext in ['.jpg','.jpeg','.png','.zip','.html']:return YT+'/90_Архив_этапов/00_Общие_сравнения/'+p.name
 return TECH+'/'+old

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def dims(p):
 if p.suffix.lower() in ['.png','.jpg','.jpeg','.webp','.gif']:
  try:
   with Image.open(p) as im:return f'{im.width} × {im.height}'
  except Exception:return ''
 if p.suffix.lower()=='.svg':
  s=p.read_text()[:600];w=re.search(r'<svg[^>]*\bwidth="([\d.]+)',s);h=re.search(r'<svg[^>]*\bheight="([\d.]+)',s)
  if w and h:return w.group(1)+' × '+h.group(1)
 return ''
def platform(old):
 if old.startswith('uploads/'):return 'Материалы от вас'
 if '/telegram_avatars/' in old or '/emoji/' in old or '/branding/telegram/' in old:return 'Telegram'
 if '/github/' in old:return 'GitHub'
 if '/branding/vk/' in old:return 'VK'
 if '/brand/' in old or '/branding/' in old and not any(x in old for x in ['/youtube/','/guides/','/streams/']):return 'Бренд / несколько площадок'
 return 'YouTube'

if PLAN.exists():
 plan=json.loads(PLAN.read_text());records=plan['files']
else:
 videos={r['id']:r for r in json.loads((ROOT/'moko/final/registry_final.json').read_text())}
 emoji={r['slug']:r['title'] for r in json.loads((ROOT/'moko/emoji/source/manifest.json').read_text())}
 avatars={r['slug']:r['title'] for r in json.loads((ROOT/'moko/telegram_avatars/final_graphic/manifest.json').read_text())}
 files=sorted([p for root in ['moko','uploads'] for p in (ROOT/root).rglob('*') if p.is_file()])
 records=[]
 for p in files:
  old=str(p.relative_to(ROOT));new=route(old);ext=p.suffix.lower();role='';status='Подготовлено';name=p.stem;theme='';videoid=''
  if new.startswith(TECH):role='Рабочий код / данные';status='Техническое'
  elif '/90_Архив' in new:role='Архив этапов';status='Архив — не текущий финал'
  elif '/81_Согласованные' in new:role='Исходник согласованного портрета';status='Согласованный образец'
  elif '/80_Исходные' in new:role='Исходная обложка YouTube';status='Исходник'
  elif new.startswith(TGF):role='Аватары Telegram';status='ФИНАЛ · исходный В'
  elif new.startswith(EM):role='Кастомные эмодзи';status='24 статичных · не опубликовано'
  elif new.startswith(GHA):role='Аватар GitHub';status='Готовый файл · не установлен'
  elif new.startswith(YT94):role='Обложки и документы 94 видео';status='18 согласовано / 76 на проверке'
  elif new.startswith(YT+'/02_'):role='Оформление YouTube'
  elif new.startswith(YT+'/03_'):role='Статичные сцены эфиров'
  elif new.startswith(VK+'/01_'):role='Обложки VK'
  elif new.startswith(VK+'/02_'):role='Меню VK — 8 плиток'
  elif new.startswith(TG+'/03_'):role='Баннер поста / необязательные обои'
  elif new.startswith(BR):role='Бренд / логотипы / базовое оформление'
  elif new.startswith(RAW):role='Предоставленный материал';status='Оригинал'
  if p.stem in videos:
   v=videos[p.stem];name=v['title'];theme=v['playlist'];videoid=v['id']
   if old.startswith('moko/production/thumbnails/') or old.startswith('moko/final/shorts/'):status=v['status']
  if old.startswith('moko/emoji/') and p.stem in emoji:name=emoji[p.stem]
  if '/telegram_avatars/' in old and p.stem in avatars:name=avatars[p.stem]
  if old.startswith('moko/github/'):name='MOKO · GitHub — '+p.stem
  records.append({'old':old,'new':new,'size':p.stat().st_size,'sha256':sha(p),'platform':platform(old),'section':role,'status':status,'name':name,'format':ext.lstrip('.').upper(),'dimensions':dims(p),'theme':theme,'video_id':videoid})
 assert len({r['new'] for r in records})==len(records)
 assert not any((DEST/r['new']).exists() for r in records)
 plan={'date':'2026-09-09','workspace_root':str(ROOT),'organized_root':'MOKO','original_files':len(records),'original_bytes':sum(r['size'] for r in records),'files':records}
 PLAN.write_text(json.dumps(plan,ensure_ascii=False,indent=2))
 # Minimal rollback/address table is written before the first move.
 with (ORG/'Старые_и_новые_пути.csv').open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.writer(f,delimiter=';');w.writerow(['Старый путь','Новый путь','Байт','SHA256'])
  for r in records:w.writerow([r['old'],'MOKO/'+r['new'],r['size'],r['sha256']])

for r in records:
 src=ROOT/r['old'];dst=DEST/r['new'];dst.parent.mkdir(parents=True,exist_ok=True)
 if src.exists():
  assert not dst.exists();shutil.move(str(src),str(dst))
 assert dst.exists() and dst.stat().st_size==r['size'] and sha(dst)==r['sha256'],r['old']
# Only now remove emptied legacy directories; no content is deleted.
for oldroot in ['moko','uploads']:
 p=ROOT/oldroot
 if p.exists():
  for d in sorted([d for d in p.rglob('*') if d.is_dir()],key=lambda x:len(x.parts),reverse=True):
   if not any(d.iterdir()):d.rmdir()
  if not any(p.iterdir()):p.rmdir()
# Intentional small copies: each portal has its own ready-to-find universal avatars.
extras=[]
for folder in [YT+'/04_Аватары',VK+'/03_Аватары']:
 for src in (DEST/BR/'04_Общие_аватары_YouTube_VK').glob('*'):
  dst=DEST/folder/src.name;dst.parent.mkdir(parents=True,exist_ok=True)
  if not dst.exists():shutil.copy2(src,dst)
  assert sha(src)==sha(dst)
  extras.append({'old':'Копия: '+str(src.relative_to(DEST)),'new':str(dst.relative_to(DEST)),'size':dst.stat().st_size,'sha256':sha(dst),'platform':'YouTube' if folder.startswith(YT) else 'VK','section':'Общий фирменный аватар','status':'Копия для удобства','name':'MOKO — '+('знак и название' if 'with_name' in dst.name else 'крупный знак'),'format':dst.suffix[1:].upper(),'dimensions':dims(dst),'theme':'','video_id':''})
# Verify all ZIP packages after relocation without changing their contents.
zips=[]
for r in records:
 if r['format']=='ZIP':
  with zipfile.ZipFile(DEST/r['new']) as z:assert z.testzip() is None
  zips.append(r['new'])
report={'original_files_moved':len(records),'all_original_sha256_unchanged':True,'original_bytes':sum(r['size'] for r in records),'deleted_original_files':0,'convenience_copies':len(extras),'zip_archives_verified':len(zips),'legacy_dirs_removed_when_empty':True}
(ORG/'Отчёт_проверки.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
(ORG/'Дополнительные_копии.json').write_text(json.dumps(extras,ensure_ascii=False,indent=2))
# A searchable Office catalog, with the main files on the first sheet.
wb=Workbook();ws=wb.active;ws.title='Главное';ws.append(['MOKO — где искать материалы','Путь относительно папки MOKO','Примечание'])
quick=[('YouTube — все 94 обложки',YT94+'/MOKO_Все_обложки.html','92 видео + 2 Shorts. Новые 76 ещё на проверке.'),('YouTube — скачать комплект 94',YT94+'/MOKO_Полный_комплект_94.zip','Ранее подготовленный ZIP сохранён без изменений.'),('Telegram — финальные аватары В',TGF+'/MOKO_Telegram_Финал_В.html','Только исходный вариант В; не В1 и не В2.'),('Telegram — ZIP финальных аватаров',TGF+'/MOKO_Telegram_Финал_В.zip','6 PNG + 6 SVG.'),('Telegram — эмодзи',EM+'/MOKO_Эмодзи_Просмотр.html','24 статичных эмодзи. Пак не опубликован.'),('Telegram — ZIP эмодзи',EM+'/MOKO_24_эмодзи_v01.zip','PNG 100 × 100, WebP, SVG, таблица привязок.'),('GitHub — PNG для установки',GHA+'/MOKO_GitHub_512.png','Для организации MOKO.'),('GitHub — комплект',GHA+'/MOKO_GitHub_Аватар.zip','PNG 512/1024 + SVG.'),('YouTube — шапка',YT+'/02_Оформление_канала/channel_banner.png','Загружать целиком, не обрезать до полосы.'),('VK — обложка',VK+'/01_Обложки/community_cover.png','Горизонтальная версия.'),('VK — мобильная обложка',VK+'/01_Обложки/mobile_cover_static.png','Вертикальная статичная версия.'),('Бренд — общее оформление',BR+'/05_Базовый_комплект_оформления/MOKO_Оформление_каналов.html','Финальные Telegram-аватары лежат отдельно.'),('Бренд — базовый ZIP',BR+'/05_Базовый_комплект_оформления/MOKO_Логотип_и_баннеры_v01.zip','Шапки, логотипы и другие материалы первого комплекта.')]
for name,path,note in quick:
 assert (DEST/path).is_file();ws.append([name,path,note]);ws.cell(ws.max_row,1).hyperlink='./'+path;ws.cell(ws.max_row,1).style='Hyperlink'
ws.append(['Все файлы — на листе «Файлы»','Поиск по названию, площадке, разделу, формату, видео и пути.','Ссылки на файлы работают при скачивании папки MOKO с сохранением структуры.'])
ws.column_dimensions['A'].width=43;ws.column_dimensions['B'].width=88;ws.column_dimensions['C'].width=65;ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
fs=wb.create_sheet('Файлы');headers=['Площадка','Раздел','Статус','Название / назначение','Формат','Размеры','Путь относительно MOKO','Открыть','Размер, КБ','Тема YouTube','ID видео'];fs.append(headers)
allrows=sorted(records+extras,key=lambda r:(r['new']))
for r in allrows:
 fs.append([r['platform'],r['section'],r['status'],r['name'],r['format'],r['dimensions'],r['new'],'Открыть',round(r['size']/1024,1),r['theme'],r['video_id']]);fs.cell(fs.max_row,8).hyperlink='./'+r['new'];fs.cell(fs.max_row,8).style='Hyperlink'
fs.auto_filter.ref=fs.dimensions;fs.freeze_panes='D2'
for i,w in enumerate([19,33,34,66,11,18,94,12,14,50,19],1):fs.column_dimensions[get_column_letter(i)].width=w
for sh in [ws,fs]:
 sh.row_dimensions[1].height=31
 for c in sh[1]:c.fill=PatternFill('solid',fgColor='17263A');c.font=Font(color='FFFFFF',bold=True,size=11);c.alignment=Alignment(vertical='center',wrap_text=True)
 for row in sh.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
  if row[0].row%2==0:
   for c in row:c.fill=PatternFill('solid',fgColor='F2F5F8')
 sh.sheet_view.showGridLines=False
wb.save(DEST/'00_Каталог_файлов.xlsx')
print(json.dumps(report,ensure_ascii=False,indent=2))
