from pathlib import Path
import json,csv
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from docx import Document
from docx.shared import Pt
R=Path('/home/user/moko');D=R/'docs';rows=json.load(open(R/'catalog/registry.json'));new=json.load(open(R/'batch03/new_items.json'));newby={a['id']:a for a in new};desc={a['id']:a for a in json.load(open(R/'catalog/descriptions_all.json'))}
for a in rows:
 a['description_available']=bool(desc.get(a['id'],{}).get('description'))
 if a['id'] in newby:
  n=newby[a['id']];a['headline']=n['headline'];a['status']='Готово · новая партия 03, проверить';a['file']='thumbnails/'+a['id']+'.jpg';a['basis']='Плейлист Python + исходная обложка + описание и таймкоды автора';a['notes']=(a['notes']+' '+n['note']).strip();a['review']='Проверить новый макет; весь ролик не просмотрен'
 if a['id'] in ['feEv04DYWbA','OhPQyqKovhA']:
  a['language']='LabVIEW';a['basis']='Описание автора / хэштег #labview и исходная обложка';a['notes']=a['notes'].replace('Не ставить Python/LabVIEW на обложке без подтверждения языка по содержанию.','').strip()+' Язык LabVIEW подтверждён описанием автора; основной плейлист практической разработки сохранён, LabVIEW добавлен как дополнительный.'
  a['extra']='; '.join(dict.fromkeys([x for x in a['extra'].split('; ') if x]+['LabVIEW — драйверы, утилиты и плагины']))
(R/'catalog/registry.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2))
wb=load_workbook(D/'MOKO_Реестр_94_видео_и_плейлисты.xlsx');ws=wb['Видео']
keys=['number','id','title','url','tab','existing','playlist','playlist_position','extra','visual','color','language','format','lesson','headline','status','review','notes','file','basis']
for a in rows:
 for ci,k in enumerate(keys,1):ws.cell(a['number']+1,ci,a.get(k,''))
ws2=wb['Готовые файлы']
for n in new:
 a=next(x for x in rows if x['id']==n['id']);p=R/'production/thumbnails'/(a['id']+'.jpg');ws2.append([a['id'],a['title'],a['status'],a['file'],p.stat().st_size,'1280 × 720'])
for t in ws2.tables.values():t.ref=ws2.dimensions
ws2.auto_filter.ref=ws2.dimensions
for rr in ws2.iter_rows(min_row=2):
 ws2.row_dimensions[rr[0].row].height=62
 for c in rr:c.alignment=Alignment(vertical='top',wrap_text=True)
ver=wb.create_sheet('Проверка партии 03');ver.append(['ID','Тема','Что подтверждено','Основание','Что не делалось'])
for n in new:ver.append([n['id'],n['headline'],n['note'],'Публичное описание и главы с таймкодами','Полный просмотр видео; интерфейс не имитировался'])
for c in ver[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='172332')
for col,w in {'A':17,'B':32,'C':90,'D':40,'E':55}.items():ver.column_dimensions[col].width=w
for rr in ver.iter_rows():
 for c in rr:c.alignment=Alignment(wrap_text=True,vertical='top')
for i in range(2,8):ver.row_dimensions[i].height=80
ver.freeze_panes='A2';ver.auto_filter.ref=ver.dimensions
wb.save(D/'MOKO_Реестр_94_видео_и_плейлисты_v1.1.xlsx')
with open(D/'MOKO_Реестр_94_видео_v1.1.csv','w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow([c.value for c in ws[1]]);w.writerows([[a.get(k,'') for k in keys] for a in rows])
# Update rules with the newly established evidence-first workflow.
rules=Document(D/'MOKO_Правила_обложек_v1.docx')
for p in rules.paragraphs:
 if 'версия 1.0' in p.text:p.text=p.text.replace('версия 1.0','версия 1.1')
rules.add_heading('13. Проверка смысла до выбора изображения',1)
for text in ['До создания каждой следующей обложки читать доступное описание автора и главы/таймкоды. Если название функции двусмысленно, не выбирать картинку по одному переводу слова. При недостатке сведений отметить уточнение, а не подменять тему догадкой.', 'Пример Stage: урок 3 посвящён типам сообщений info, error, plugin, driver, report, warning. Ранее предложенное «этапы программы» отклонено после проверки описания.', 'Пример Messenger: урок 4 посвящён диалоговым окнам, таймеру, вводу string/boolean, выбору файла и картинкам. Не использовать образ социального мессенджера или Telegram.', 'Пример Driver: урок 6 из курса Python — set/get и обмен с прибором. Разработка драйвера в LabVIEW — другая тема и другой плейлист.', 'Наличие описания/таймкодов не означает полного просмотра ролика. В реестре разделять «описание получено», «содержание проверено по описанию» и «ролик просмотрен».']:
 rules.add_paragraph(text)
rules.save(D/'MOKO_Правила_обложек_v1.1.docx')
summary=Document(D/'MOKO_Сводная_карта_плейлистов_v1.docx')
for p in summary.paragraphs:
 if 'Рабочая сводка v1' in p.text:p.text=p.text.replace('Рабочая сводка v1','Рабочая сводка v1.1')
 if p.text.startswith('Статус на момент выдачи:'):
  p.text='Статус на момент выдачи: 18 JPG подготовлено (6 согласованных образцов и 12 новых), 74 обычные обложки в очереди, 2 Shorts требуют отдельного вертикального решения. Ничего на YouTube не загружалось.'
summary.add_heading('Дополнение v1.1: проверка описаний и партия 03',1)
summary.add_paragraph('Получены непустые описания для 90 из 94 роликов. Это сбор источников, а не заявление о просмотре всех видео. Для шести новых уроков партии 03 описания и главы изучены перед подготовкой обложек.')
for n in new:summary.add_paragraph(n['id']+' — '+n['headline']+'. '+n['note'])
summary.add_paragraph('В описаниях стримов OhPQyqKovhA и feEv04DYWbA подтверждён LabVIEW. Их основной плейлист практической разработки сохранён; LabVIEW назначен дополнительным. В основном учебном LabVIEW-плейлисте по-прежнему 13 уроков; при включении этих двух стримов как дополнительных материалов будет 15 записей.')
summary.add_paragraph('Согласованные портреты и первые шесть обложек в этой партии не изменялись. Новые макеты имеют статус «проверить», а не «утверждено» или «загружено».')
summary.save(D/'MOKO_Сводная_карта_плейлистов_v1.1.docx')
(R/'batch03/research/review_notes.md').write_text('# Проверка партии 03\n\n'+'\n\n'.join('## '+n['id']+' — '+n['headline']+'\n'+n['note'] for n in new)+'\n\nВсе выводы — по описаниям автора и главам, не по полному просмотру. Иллюстрации условные, не имитация реального интерфейса.')
print('Registry, rules and summary updated to 1.1. Prepared files:',sum(bool(a['file']) for a in rows))
