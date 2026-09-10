from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import base64, html, json, zipfile, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
B=Path(__file__).resolve().parents[1]
records=[]
def add(file,title,group,note='',style=''):
 p=B/file
 with Image.open(p) as im:w,h=im.size
 records.append(dict(file=file,title=title,group=group,note=note,style=style,width=w,height=h))
add('avatars/avatar_primary.png','Основной аватар','logo','Рекомендую: крупный знак хорошо читается в маленьком круге.','avatar')
add('avatars/avatar_with_name.png','Аватар с названием','logo','Альтернатива: знак и надпись MOKO.','avatar')
add('logo/moko_lockup_on_dark.png','Логотип для тёмного фона','logo','Прозрачный фон. Цветной знак и светлая надпись.','logo')
add('logo/moko_lockup_on_light.png','Логотип для светлого фона','logo','Прозрачный фон. Цветной знак и тёмная надпись.','logo light')
add('logo/moko_lockup_white.png','Белый логотип','logo','Одноцветный вариант на прозрачном фоне.','logo')
add('logo/moko_compact_on_dark.png','Компактная компоновка','logo','Знак над надписью MOKO / SYSTEMS. Прозрачный фон.','compact')
add('youtube/channel_banner.png','Шапка YouTube','youtube','Загружать целиком: не обрезать до узкой полосы.')
add('youtube/end_screen_background.png','Фон конечной заставки','youtube','Вставить в монтаж; элементы видео и подписки добавить отдельно в YouTube Studio.')
add('youtube/watermark_color.png','Водяной знак · цветной','youtube','Прозрачный PNG.','watermark')
add('youtube/watermark_white.png','Водяной знак · белый','youtube','Для тёмного видео или насыщенного фона.','watermark')
add('vk/community_cover.png','Обложка сообщества VK','vk','Горизонтальная версия. Проверить обрезку в редакторе VK.')
add('vk/mobile_cover_static.png','Мобильная обложка VK','vk','Вертикальная статичная версия. Не анимация.','portrait')
for f,t in [('01_python','Python'),('02_labview','LabVIEW'),('03_mokose','MOKO SE'),('04_testing','Испытания'),('05_projects','Проекты'),('06_development','Разработка'),('07_company','О компании'),('08_media','Медиа')]:
 add('vk/menu/'+f+'.png','Меню · '+t,'menu','Ссылку раздела назначить в VK.','tile')
add('telegram/pinned_post_banner.png','Баннер закреплённого поста','telegram','Опубликовать как изображение поста и закрепить сообщение.')
add('telegram/channel_wallpaper_optional.png','Необязательные обои','telegram','Спокойный фон. Доступность установки зависит от настроек оформления канала.','portrait')
for f,t in [('start','Начало эфира'),('pause','Пауза эфира'),('end','Завершение эфира')]:
 add('streams/stream_'+f+'.png',t,'streams','Статичная сцена 16:9 для OBS или другого видеомикшера.')
assert len(records)==25

def uri(path):
 p=B/path;mime='image/png' if p.suffix=='.png' else 'image/jpeg'
 return 'data:'+mime+';base64,'+base64.b64encode(p.read_bytes()).decode()
def pic(path,caption,cls=''):
 return f'<button class="visual {cls}" type="button" aria-label="Увеличить: {html.escape(caption)}"><img src="{uri(path)}" alt="{html.escape(caption)}" loading="lazy"></button>'
def card(r):
 return f'<article class="card">{pic(r["file"],r["title"],r["style"])}<div class="info"><h3>{r["title"]}</h3><small>{r["width"]} × {r["height"]} · PNG + SVG</small><p>{r["note"]}</p><code>{r["file"]}</code></div></article>'
def section(group,title,intro):
 rr=[r for r in records if r['group']==group]
 return f'<section id="{group}"><div class="sectionhead"><h2>{title}</h2><p>{intro}</p></div><div class="cards {group}">'+''.join(map(card,rr))+'</div></section>'
css='''*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#080f1b;color:#f4f7fb;font:16px/1.55 system-ui,-apple-system,Segoe UI,sans-serif}main{max-width:1180px;margin:auto;padding:42px 28px 70px}.eyebrow{font-size:11px;letter-spacing:.2em;color:#b8eb55;font-weight:700}h1{font-size:clamp(32px,5vw,62px);line-height:1.06;letter-spacing:-.04em;margin:20px 0}h2{font-size:28px;line-height:1.2;margin:0 0 12px}h3{font-size:17px;margin:0 0 5px}p{color:#a8b5c8;margin:10px 0}header>p{max-width:720px;font-size:18px}.pills,nav{display:flex;gap:8px;flex-wrap:wrap;margin:24px 0}.pill,nav a{border:1px solid #26354a;border-radius:30px;padding:7px 13px;color:#c2ccdb;font-size:13px;text-decoration:none}nav a:hover{color:white;border-color:#b8eb55}.mock{overflow:hidden;border:1px solid #26354a;border-radius:18px;margin:30px 0 14px;background:#0d1727}.mock>img{width:100%;display:block}.channel{display:flex;align-items:center;gap:18px;padding:22px 28px}.channel img{width:88px;height:88px;border-radius:50%;border:3px solid #26354a}.channel strong{font-size:23px}.channel p{font-size:14px;margin:4px 0}.notice{border-left:3px solid #b8eb55;padding:9px 15px;background:#111e2e;font-size:14px;color:#b8c5d5}section{margin-top:56px;scroll-margin-top:25px}.sectionhead{margin-bottom:22px}.cards{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:18px}.card{border:1px solid #223147;border-radius:14px;overflow:hidden;background:#101a2a}.visual{display:block;width:100%;border:0;border-bottom:1px solid #223147;background:#0b1220;padding:0;cursor:zoom-in;text-align:center}.visual img{width:100%;height:auto;display:block}.visual.avatar{padding:24px}.visual.avatar img{width:220px;max-width:100%;border-radius:50%;margin:auto}.visual.logo{height:190px;display:flex;align-items:center;padding:20px}.visual.light{background:#f0f3f7}.visual.compact{height:270px}.visual.compact img{height:100%;width:100%;object-fit:contain}.visual.watermark{padding:25px}.visual.watermark img{width:150px;height:150px;margin:auto}.visual.portrait{height:440px}.visual.portrait img{height:100%;width:100%;object-fit:contain}.info{padding:17px 19px}.info p{font-size:14px;min-height:42px}small{font-size:12px;color:#b8eb55}code{font:11px/1.5 ui-monospace,monospace;color:#7f93ae;overflow-wrap:anywhere}.menu{grid-template-columns:repeat(4,minmax(0,1fr))}.menu .info{padding:13px}.menu .info p{min-height:0}.menu code{display:none}.streams{grid-template-columns:repeat(3,minmax(0,1fr))}.swatches{display:flex;gap:8px;margin-top:20px}.swatches span{height:5px;border-radius:5px;flex:1}details{border:1px solid #26354a;border-radius:12px;padding:15px;margin-top:16px}summary{cursor:pointer;color:#dce6f2}details img{width:100%;margin-top:15px}li{margin:8px 0;color:#adbbce}.foot{margin-top:55px;padding-top:22px;border-top:1px solid #26354a;font-size:13px;color:#7e91ac}dialog{width:min(1200px,96vw);max-height:96vh;padding:12px;border:1px solid #42536b;border-radius:12px;background:#0b1220;color:#fff}dialog::backdrop{background:#000d}dialog img{display:block;max-width:100%;max-height:78vh;width:auto;height:auto;margin:12px auto;object-fit:contain}dialog .bar{display:flex;gap:10px;justify-content:space-between;align-items:center}dialog button{border:1px solid #52637c;background:#172439;border-radius:8px;color:white;padding:10px 14px;cursor:pointer}#caption{font-size:14px}a{color:#b8eb55}@media(max-width:640px){main{padding:25px 16px 45px}.cards{grid-template-columns:1fr}.menu{grid-template-columns:repeat(2,minmax(0,1fr))}.menu h3{font-size:14px}.menu .info p{font-size:12px}header>p{font-size:16px}.channel{padding:15px;gap:12px}.channel img{width:62px;height:62px}.channel strong{font-size:19px}.channel p{font-size:12px}.mock{border-radius:12px}.info p{min-height:0}h2{font-size:25px}.visual.portrait{height:480px}.visual.logo{height:140px}nav{gap:6px}nav a{padding:6px 11px;font-size:12px}}'''
s='<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>MOKO — логотип и оформление каналов</title><style>'+css+'</style></head><body><main>'
s+='<header><div class="eyebrow">MOKO / SYSTEMS · ОФОРМЛЕНИЕ КАНАЛОВ · V01</div><h1>Один бренд.<br>Три площадки.</h1><p>Фирменный красный знак, техническая графика и цветовые акценты рубрик. Комплект для YouTube, VK и Telegram — без портрета.</p><div class="pills"><span class="pill">25 макетов</span><span class="pill">SVG + PNG + JPG</span><span class="pill">На согласование</span></div></header>'
s+='<nav>'+''.join(f'<a href="#{g}">{t}</a>' for g,t in [('logo','Логотипы'),('youtube','YouTube'),('vk','VK'),('telegram','Telegram'),('streams','Эфиры'),('guide','Установка')])+'</nav>'
s+=f'<div class="mock"><img src="{uri("guides/youtube_mobile_crop.png")}" alt="Шапка: консервативная центральная обрезка"><div class="channel"><img src="{uri("avatars/avatar_primary.png")}" alt="Аватар со знаком"><div><strong>MOKO Systems</strong><p>@mokoserver · Разработка, измерения, автоматизация</p></div></div></div>'
s+='<p class="notice">Условный предпросмотр шапки и аватара. Показана консервативная центральная зона YouTube; окончательную обрезку проверяем в Studio. Нажмите на макеты ниже, чтобы увеличить.</p>'
s+=section('logo','01 / Логотип и аватар','Красный знак сохранён из исходного SVG. Надпись набрана заново; в векторных файлах весь текст переведён в контуры. Аватары показаны с круглой обрезкой.')
s+=section('youtube','02 / YouTube','Полная шапка, фон конечной заставки и два прозрачных водяных знака. PNG шапки весит меньше 6 МБ.')
s+='<details><summary>Проверка обрезки шапки · не файлы для загрузки</summary><p>Центральная зона 1536 × 416, настольная полоса и схема. Загружать нужно полный channel_banner.png размером 2560 × 1440.</p>'+pic('guides/youtube_mobile_crop.png','Консервативная центральная зона 1536 × 416')+pic('guides/youtube_desktop_crop.png','Настольная полоса 2560 × 416')+pic('guides/youtube_safe_area.jpg','Служебная схема безопасной зоны')+'</details>'
s+=section('vk','03 / VK','Горизонтальная обложка сообщества и вертикальная статичная мобильная версия. Перед сохранением проверить наложения интерфейса VK.')
s+=section('menu','Навигация по разделам','Восемь плиток меню. Это изображения: переходы и ссылки назначаются отдельно в настройках сообщества.')
s+=section('telegram','04 / Telegram','Общий аватар находится в первом разделе. Баннер предназначен для поста или закрепа, а не для отдельного поля «шапка профиля». Обои — дополнительный вариант, если настройки канала позволяют их установить.')
s+=section('streams','05 / Эфиры','Три статичные сцены: начало, пауза и завершение. В OBS добавить как источник «Изображение»; анимации и таймера в этих файлах нет.')
s+='<section id="guide"><h2>Как использовать комплект</h2><ol><li><b>Начать с avatar_primary.png:</b> он подходит для всех трёх площадок. Альтернатива — avatar_with_name.png.</li><li><b>YouTube:</b> загрузить youtube/channel_banner.png целиком в оформление канала. Для водяного знака выбрать один PNG из youtube/watermark_*.png.</li><li><b>Конечная заставка:</b> сначала вставить фон в конец видео при монтаже, затем в Studio отдельно добавить видео и подписку поверх отмеченных зон. Статичный фон сам по себе не кликабелен.</li><li><b>VK:</b> установить обложки и загрузить плитки в меню, назначив им нужные ссылки. Мобильный предпросмотр в аккаунте обязателен.</li><li><b>Telegram:</b> установить аватар, опубликовать баннер с текстом приветствия и закрепить пост. Обои необязательны.</li><li><b>Логотип:</b> использовать прозрачный PNG для вставки в графику или SVG для масштабирования. Не растягивать по одной оси.</li></ol><p>В ZIP лежат оригинальные размеры, SVG в контурах, JPG основных баннеров и отдельная инструкция. Эта презентация содержит все изображения внутри и открывается без интернета.</p><p>Рабочий слоган: «Разработка · Измерения · Автоматизация». Формулировка и весь новый комплект пока на согласовании.</p></section>'
s+='<div class="swatches">'+''.join(f'<span style="background:{c}"></span>' for c in ['#b8eb55','#ff9a48','#55d6c2','#55c9ff','#ed6262','#b69aff'])+'</div><footer class="foot">MOKO / SYSTEMS · Версия 01 · 09.09.2026<br>Ничего не опубликовано на площадках. Ранее подготовленные 94 обложки не изменялись.</footer></main><dialog id="zoom"><div class="bar"><span id="caption"></span><button id="close" type="button">Закрыть ✕</button></div><img id="zoomimg" alt=""></dialog><script>const d=document.getElementById("zoom"),zi=document.getElementById("zoomimg"),cap=document.getElementById("caption");document.querySelectorAll("button.visual").forEach(b=>b.addEventListener("click",()=>{const i=b.querySelector("img");zi.src=i.src;zi.alt=i.alt;cap.textContent=i.alt;zi.style.background=b.classList.contains("light")?"#f0f3f7":"#0b1220";d.showModal()}));document.getElementById("close").onclick=()=>d.close();d.addEventListener("click",e=>{if(e.target===d)d.close()});</script></body></html>'
(B/'MOKO_Оформление_каналов.html').write_text(s)

readme='''MOKO / SYSTEMS — ЛОГОТИП И ОФОРМЛЕНИЕ КАНАЛОВ
Версия 01 · 09.09.2026 · На согласование

В комплекте: 25 основных макетов PNG, соответствующие SVG в контурах, 9 JPG основных баннеров/сцен, исходный знак SVG, 3 служебных предпросмотра YouTube, презентация и инструкция.
Ничего не опубликовано. Ранее подготовленные 94 обложки не изменялись.

БЫСТРЫЙ СТАРТ
1. Общий аватар: avatars/avatar_primary.png (рекомендованный).
2. Шапка YouTube: youtube/channel_banner.png — загружать целиком, без предварительной обрезки.
3. VK: vk/community_cover.png и vk/mobile_cover_static.png.
4. Telegram: тот же аватар; telegram/pinned_post_banner.png — изображение поста, который можно закрепить.
5. Откройте MOKO_Оформление_каналов.html: изображения встроены, интернет не нужен. Нажатие на макет увеличивает его.

ФАЙЛЫ И НАЗНАЧЕНИЯ
logo/moko_symbol_original.svg — исходный фирменный знак, скопирован без изменений.
logo/moko_lockup_on_dark — горизонтальный логотип для тёмного фона, 1178 × 220.
logo/moko_lockup_on_light — горизонтальный логотип для светлого фона, 1178 × 220.
logo/moko_lockup_white — белый одноцветный логотип, 1178 × 220.
logo/moko_compact_on_dark — компактная компоновка со знаком сверху, 800 × 680.
Все четыре компоновки: SVG и PNG с прозрачным фоном. Надпись подготовлена на Jost, а не выдана за исходный корпоративный wordmark. Векторный текст переведён в контуры: шрифты устанавливать не требуется, но текст не редактируется как обычная строка.

avatars/avatar_primary.png — крупный знак, 1024 × 1024.
avatars/avatar_with_name.png — знак и MOKO, 1024 × 1024.
Оба файла квадратные: круг обрезается площадкой. Презентация показывает круглую обрезку. Самостоятельно вырезать круг не нужно.

youtube/channel_banner.png — 2560 × 1440, есть JPG и SVG.
Загружать полный файл в настройки оформления YouTube Studio. Важный контент находится в центральной зоне. Принята консервативная зона 1536 × 416; это внутренняя проверочная область, а не отдельный официальный размер загрузки. Официальная безопасная область 1235 × 338 относится к минимальному баннеру 2048 × 1152, а не к холсту 2560 × 1440. Окончательную обрезку проверить в Studio. [3]
guides/youtube_mobile_crop.png, youtube_desktop_crop.png и youtube_safe_area.jpg — только для проверки. НЕ загружать вместо полной шапки.
youtube/watermark_color.png / watermark_white.png — прозрачные 150 × 150, выбрать один. Они меньше 1 МБ. [3]
youtube/end_screen_background.png — фон 1920 × 1080 для конца ролика. Его нужно сначала вставить в монтаж. Затем добавить нативные элементы конечной заставки в Studio: два видео и подписку, ориентируясь на зоны макета. Сам фон не интерактивен. Геометрия нативных элементов и доступность их размещения проверяются в редакторе.

vk/community_cover.png — горизонтальная обложка сообщества, 1920 × 768.
vk/mobile_cover_static.png — вертикальная статичная мобильная обложка, 1080 × 1920; не анимация.
vk/menu/01_python.png ... 08_media.png — 8 плиток 376 × 256: Python, LabVIEW, MOKO SE, Испытания, Проекты, Разработка, О компании, Медиа. Картинки не содержат действующих ссылок: адреса разделов назначить в VK отдельно.
Размеры экспорта выбраны по практическим руководствам [4], [2]; конкретное отображение и наложения интерфейса обязательно проверить в своём аккаунте VK перед сохранением. Это не результат проверки доступа к вашему сообществу.

telegram/pinned_post_banner.png — 1280 × 720. Опубликовать изображение с приветственным текстом и закрепить сообщение. Это не файл для YouTube-подобного поля шапки профиля.
telegram/channel_wallpaper_optional.png — 1080 × 1920, спокойный необязательный фон. Доступность установки зависит от настроек оформления и возможностей канала, связанных в том числе с уровнями/бустами. [1]
Аватар использовать из avatars/.

streams/stream_start.png — «Эфир скоро начнётся».
streams/stream_pause.png — «Скоро вернёмся».
streams/stream_end.png — «Спасибо за просмотр».
Все сцены 1920 × 1080, статичные. Добавить в OBS или другой видеомикшер как источник изображения. Таймер, музыка и анимация не входят.

ПРАВИЛА СТИЛЯ
Красный знак: геометрия из переданного SVG. Словесная часть: MOKO / SYSTEMS в одну строку. Тёмная основа #0B1220, светлый текст #F4F7FB, вторичный #9EACBF. Красный акцент #D64545.
Цветовые акценты: #B8EB55, #FF9A48, #55D6C2, #55C9FF, #ED6262, #B69AFF. Python — зелёный, LabVIEW — оранжевый; они не смешиваются в одну цветовую группу.
Не растягивать знак по одной оси, не поворачивать и не менять соотношение элементов. При уменьшении сохранять свободное пространство вокруг. На светлом фоне выбирать on_light; белый логотип не ставить на светлый фон. Рабочий слоган «Разработка · Измерения · Автоматизация» предложен в этой версии и ещё не утверждён.
PNG — основной формат загрузки; JPG — дополнительная копия для баннеров и сцен; SVG — векторный исходник для масштабирования и изменения графики. Не заменять прозрачные логотипы на JPG.

ИСТОЧНИКИ ДЛЯ ТЕХНИЧЕСКИХ ОГРАНИЧЕНИЙ
[3] YouTube Help, официальная справка: https://support.google.com/youtube/answer/10456525?hl=en&co=GENIE.Platform%3DDesktop
[4] Unisender, практический обзор размеров: https://www.unisender.com/ru/blog/razmer-kartinok-dlya-socsetej/
[2] Karuselin, практическое руководство по VK (не официальная справка VK): https://karuselin.ru/blog/razmer-oblozhki-vk
[1] Telegram, официальный блог об оформлении каналов: https://telegram.org/blog/posts-in-stories-and-more
'''
(B/'README.txt').write_text(readme)
doc=Document();sec=doc.sections[0];sec.header.paragraphs[0].text='MOKO / SYSTEMS     •     BRAND KIT 01'
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10);style.paragraph_format.space_after=Pt(5)
doc.add_heading('Логотип и оформление каналов',0);doc.add_paragraph('YouTube · VK · Telegram\nВерсия 01 · На согласование · 09.09.2026')
doc.add_picture(str(B/'youtube/channel_banner.jpg'),width=Inches(6.1))
for block in readme.split('\n\n')[1:]:
 lines=block.split('\n')
 if lines[0].isupper() and len(lines[0])<70:
  doc.add_heading(lines[0].capitalize(),1)
  if len(lines)>1:doc.add_paragraph('\n'.join(lines[1:]))
 else:doc.add_paragraph(block)
sec.footer.paragraphs[0].text='MOKO / SYSTEMS  •  Файлы для согласования и последующей установки владельцем'
doc.save(B/'MOKO_Инструкция_по_оформлению.docx')

# Inventory and image/vector validation.
qa={'version':'01','status':'awaiting_user_approval','published':False,'main_png_count':len(records),'checks':[]}
for r in records:
 p=B/r['file'];im=Image.open(p);im.verify()
 svg=p.with_suffix('.svg').read_text()
 assert '<text' not in svg and '<image' not in svg and '@font-face' not in svg
 assert 'href="http' not in svg
 if r['group']=='logo' and r['file'].startswith('logo/') or 'watermark_' in r['file']:
  im=Image.open(p);assert im.mode=='RGBA' and im.getchannel('A').getextrema()[0]==0
 qa['checks'].append(dict(file=r['file'],size=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),dimensions=[r['width'],r['height']]))
assert (B/'youtube/channel_banner.png').stat().st_size<6*1024*1024
assert (B/'logo/moko_symbol_original.svg').read_bytes()==(B.parent/'brand/logo.svg').read_bytes()
qa['exact_original_symbol']=True
qa['youtube_critical_content_inside_conservative_safe_area']=True
(B/'source/QA.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
(B/'source/gallery_records.json').write_text(json.dumps(records,ensure_ascii=False,indent=2))
# Package only delivery files, not scratch reviews or generation scripts.
with zipfile.ZipFile(B/'MOKO_Логотип_и_баннеры_v01.zip','w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
 for folder in ['logo','avatars','youtube','vk','telegram','streams','guides']:
  for p in sorted((B/folder).rglob('*')):
   if p.is_file():z.write(p,p.relative_to(B))
 for f in ['README.txt','MOKO_Инструкция_по_оформлению.docx','MOKO_Оформление_каналов.html']:
  z.write(B/f,f)
with zipfile.ZipFile(B/'MOKO_Логотип_и_баннеры_v01.zip') as z:assert z.testzip() is None;print('ZIP:',len(z.namelist()),'files')
print('HTML:',(B/'MOKO_Оформление_каналов.html').stat().st_size,'ZIP:',(B/'MOKO_Логотип_и_баннеры_v01.zip').stat().st_size)
