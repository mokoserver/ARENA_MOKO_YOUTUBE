from pathlib import Path
import json,hashlib,csv,collections
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path('/home/user/moko');F=R/'final';D=F/'docs'
rows=json.load(open(R/'catalog/registry.json'));groups=json.load(open(R/'catalog/proposed_playlists.json'));new=json.load(open(F/'briefs/new_76.json'));newby={a['id']:a for a in new}
low=[2,3,11,14,42,44,49,52,57,63,64,67,69,83,84,85,88,89,90,91,92]
orders={k:sorted(g['nums'],key=lambda n:rows[n-1].get('playlist_position',n)) for k,g in groups.items()}
orders['PY']=sorted(groups['PY']['nums'],key=lambda n:tuple(map(int,rows[n-1]['lesson'].split('.'))))
orders['LV']=[21,20,19,18,17,81,65,51,70,79,80,78,77]
orders['SE']=[50,45,46,7,8,54,82,53,75,66,48,74,76,59,58,47,13,72,62]
for k,nums in orders.items():
 for pos,n in enumerate(nums,1):rows[n-1]['playlist_position']=pos
for a in rows:
 n=a['number'];nb=newby.get(a['id']);previous=not nb
 if nb:
  a['headline']=nb['headline'];a['file']=nb['file'];a['basis']=nb['basis'];a['visual_source']=nb['note'];a['notes']=(a['notes']+' '+nb['note']).strip()
  a['status']='Готово · проверить новый макет' if a['tab']!='Shorts' else 'Готово · вертикальный макет, применение отдельно'
  a['review']='Финальное визуальное согласование'
 else:
  a['status']='Утверждено · прежние 18 сохранены';a['review']='';a['visual_source']='Ранее согласованный образец/партия. Файл в этой итерации не менялся.'
 if n in low:
  a['review']+='; ограничение качества исходного фото';a['notes']+=' Чёткость фото ограничена исходной обложкой. Для повышения качества нужен оригинальный кадр/снимок; оборудование не дорисовывалось.'
 if n in [7,8]:a['review']+='; выбрать оригинал или перезалив для плейлиста'
 if n in [21,62,73]:a['review']+='; архивная версия программы/ОС'
 if n in [93,94]:a['review']+='; не обычная пользовательская обложка Shorts'
 if n==94:
  a['extra']='; '.join(dict.fromkeys([x for x in a['extra'].split('; ') if x]+[groups['LV']['name']]))
 p=(F/'shorts'/(a['id']+'.jpg')) if a['tab']=='Shorts' else (R/'production/thumbnails'/(a['id']+'.jpg'))
 a['resolution']=' × '.join(map(str,Image.open(p).size));a['bytes']=p.stat().st_size;a['sha256']=hashlib.sha256(p.read_bytes()).hexdigest()
 expected=(1080,1920) if a['tab']=='Shorts' else (1280,720);assert Image.open(p).size==expected;assert p.stat().st_size<2_000_000
assert len(rows)==94 and len(set(a['id'] for a in rows))==94
(F/'registry_final.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2));(R/'catalog/registry.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2))
# Final workbook: one authoritative row for every video.
wb=Workbook();ws=wb.active;ws.title='Все 94 видео'
cols=[('number','№'),('id','YouTube ID'),('title','Исходное название'),('url','YouTube'),('playlist','Основной плейлист'),('playlist_position','Порядок'),('extra','Дополнительные плейлисты'),('visual','Визуальная рубрика'),('color','HEX'),('language','Язык / среда'),('format','Формат'),('lesson','Урок / часть'),('headline','Надпись на обложке'),('status','Статус'),('file','Файл в архиве'),('resolution','Разрешение'),('bytes','Размер, байт'),('review','Проверка перед загрузкой'),('basis','Основание тематики'),('visual_source','Источник визуала'),('notes','Примечания'),('existing','Старые плейлисты'),('sha256','SHA-256')]
ws.append([t for k,t in cols])
for a in rows:ws.append([a.get(k,'') for k,t in cols])
for rr in ws.iter_rows(min_row=2):
 rr[3].hyperlink=rr[3].value;rr[3].style='Hyperlink';rr[8].fill=PatternFill('solid',fgColor=rr[8].value.lstrip('#'));rr[8].font=Font(color='172332')
pws=wb.create_sheet('Плейлисты');pws.append(['Код','Название','Основных видео','Описание','Порядок','Цвет'])
for k,g in groups.items():pws.append([k,g['name'],len(g['nums']),g['description'],g['sorting'],g['color']])
extras=wb.create_sheet('Пересечения плейлистов');extras.append(['ID','Исходное название','Основной плейлист','Дополнительный плейлист','URL'])
for a in rows:
 for extra in a['extra'].split('; '):
  if extra:extras.append([a['id'],a['title'],a['playlist'],extra,a['url']])
check=wb.create_sheet('Проверить отдельно');check.append(['ID','Ситуация','Рекомендуемое действие'])
for a in rows:
 if a['number'] in low+[7,8,21,62,73,93,94]:check.append([a['id'],a['review'],a['notes']])
for vals in [('7ayoYanUV-w','Нет названия; отсутствует в публичной ленте','Проверить доступность; обложка не изготавливалась'),('Y9rNN698aM8','USB with LabVIEW: только в старом плейлисте','Уточнить автора/принадлежность. Не менять чужую обложку'),('_FvfYFL_1vM','Tibo 2021: только в старом плейлисте','Вероятно внешняя публикация; уточнить принадлежность')]:check.append(vals)
oldws=wb.create_sheet('Старые плейлисты');oldws.append(['Старое название','URL','Как использовать'])
old=json.load(open(R/'catalog/existing_playlists.json'))
changes=['Сохранить узкий курс и добавить в основной LabVIEW-плейлист; одна позиция требует проверки.','Сохранить как последовательный Python-курс, упорядочить по номерам.','Оставить необязательной подборкой новых публикаций.','Разнести продуктовые анонсы, новости компании и медиапортфолио.','Разделить на LabVIEW и общую разработку.','Разнести между испытаниями и электроникой/оборудованием.','Сохранить как возможности и инструкции MOKO SE, отдельно от курсов.','Оставить дополнительной подборкой внешних публикаций о компании.','События и интервью — в компанию, демонстрации платформы — в MOKO SE.']
for p,c in zip(old,changes):oldws.append([p['title'],'https://www.youtube.com/playlist?list='+p['id'],c])
qa=wb.create_sheet('Контроль выпуска');qa.append(['Проверка','Результат']);
for x in [('Всего уникальных видео',94),('Горизонтальные JPG 1280×720',92),('Вертикальные JPG 1080×1920',2),('Прежние обложки сохранены',18),('Новые обычные макеты',74),('Новые вертикальные макеты',2),('Все JPG менее 2 МБ','Да'),('Имена файлов совпадают с ID','Да'),('Полный просмотр всех видео','Нет; использованы названия, обложки, описания и главы'),('Загрузка на YouTube','Не выполнялась'),('Реальные приборы/интерфейсы перерисованы','Нет в новом массовом выпуске; ранее согласованные генеративные концепты сохранены'),('Дубликат стрима','Оба файла есть; выбор основной версии делает владелец')]:qa.append(x)
for si,s in enumerate(wb.worksheets):
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions;s.sheet_view.showGridLines=False
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='172332');c.alignment=Alignment(wrap_text=True,vertical='center')
 s.row_dimensions[1].height=36
 for rr in s.iter_rows(min_row=2):
  s.row_dimensions[rr[0].row].height=65
  for c in rr:c.alignment=Alignment(wrap_text=True,vertical='top')
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=30
 t=Table(displayName='Final'+str(si),ref=s.dimensions);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);s.add_table(t)
widths={'A':6,'B':16,'C':65,'D':44,'E':52,'F':12,'G':56,'H':31,'I':13,'J':25,'K':23,'L':13,'M':40,'N':42,'O':30,'P':20,'Q':17,'R':57,'S':61,'T':70,'U':90,'V':55,'W':70}
for c,w in widths.items():ws.column_dimensions[c].width=w
pws.column_dimensions['B'].width=58;pws.column_dimensions['D'].width=77;pws.column_dimensions['E'].width=88;extras.column_dimensions['B'].width=70;extras.column_dimensions['D'].width=60;check.column_dimensions['C'].width=105;oldws.column_dimensions['C'].width=95;qa.column_dimensions['A'].width=55;qa.column_dimensions['B'].width=100
wb.save(D/'MOKO_Реестр_94_видео_финальный.xlsx')
with open(D/'MOKO_Реестр_94_видео.csv','w',encoding='utf-8-sig',newline='') as fh:
 w=csv.writer(fh);w.writerow([t for k,t in cols]);w.writerows([[a.get(k,'') for k,t in cols] for a in rows])
# Reusable Word formatting and active video links.
def document(title,subtitle):
 d=Document();sec=d.sections[0];sec.left_margin=sec.right_margin=Cm(2);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7)
 st=d.styles['Normal'];st.font.name='Calibri';st.font.size=Pt(10);st.paragraph_format.space_after=Pt(5)
 for s in ['Title','Heading 1','Heading 2']:d.styles[s].font.name='Calibri';d.styles[s].font.color.rgb=RGBColor.from_string('172332')
 sec.header.paragraphs[0].text='MOKO / SYSTEMS  ·  YOUTUBE'
 fp=sec.footer.paragraphs[0];fp.alignment=2;fp.add_run('09.09.2026 · ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
 d.add_heading(title,0);d.add_paragraph(subtitle,'Subtitle');return d

def link(d,txt,url):
 p=d.add_paragraph();p.add_run(txt+' ');h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'27669B');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text='YouTube ↗';r.append(t);h.append(r);p._p.append(h)
def table(d,heads,data):
 t=d.add_table(rows=1,cols=len(heads));t.style='Light Shading Accent 1'
 for c,s in zip(t.rows[0].cells,heads):c.text=str(s)
 for rr in data:
  for c,s in zip(t.add_row().cells,rr):c.text=str(s)
 return t
summary=document('Итоговая карта плейлистов','MOKO Systems · полный комплект для 94 доступных роликов\n92 обычные обложки + 2 вертикальных макета Shorts')
summary.add_heading('1. Что входит в выпуск',1)
for t in ['Подготовлены файлы для всех 94 доступных роликов из собранного публичного списка: 92 JPG 1280×720 и 2 JPG 1080×1920. Прежние 18 макетов сохранены; новые 76 необходимо окончательно посмотреть перед применением.', 'Все файлы названы по YouTube ID. Реестр Excel связывает ID, исходное название, файл, основной плейлист, дополнительные плейлисты, язык, цвет и статус.', 'Проверка содержания проводилась по названиям, текущим плейлистам, исходным обложкам, доступным описаниям автора и главам. Полный просмотр всех роликов не выполнялся. Старые низкокачественные фотографии не заменялись выдуманным оборудованием.', 'Три позиции из старых плейлистов отсутствуют в публичной ленте; они не входят в 94 и вынесены на отдельную проверку. Изменения в YouTube Studio не выполнялись.']:summary.add_paragraph(t)
summary.add_heading('2. Основные разделы',1)
table(summary,['Плейлист','Видео','Визуальный акцент'],[[g['name'],len(g['nums']),g['color']] for k,g in groups.items()])
summary.add_paragraph('Сумма основных назначений — 94. Каждый ролик имеет ровно один основной плейлист; дополнительные подборки могут пересекаться. «Стрим» — формат, а не отдельный цвет.')
summary.add_heading('3. Python и LabVIEW',1)
summary.add_paragraph('Python — 21 урок, зелёный #B8EB55, порядок по номеру урока. LabVIEW — 13 основных уроков, оранжевый #FF9A48. В LabVIEW дополнительно можно включить два профильных стрима (4.1 и 4.2) и Short про Rozum Robotics — всего 16 позиций при добавлении этих практических материалов после курса.')
summary.add_paragraph('Название функции не определяет язык. Driver, Plugin и Utility встречаются внутри курса Python. Видеоуроки LabVIEW определены по названию, исходной обложке/серии и описанию, а не по одному слову «драйвер».')
summary.add_heading('4. Уточнения после изучения описаний',1)
for t in ['Stage в уроке Python №3 — типы сообщений, не этапы программы. Messenger №4 — диалоговые окна и ввод данных, не Telegram.', 'Оригинал и перезалив стрима №3 об автообновлении перенесены в «MOKO SE — возможности и инструкции». Формат остаётся стримом.', 'Объединение отчётов на примере поверки манометров находится прежде всего в инструкциях MOKO SE; испытания/поверка — дополнительная подборка.', 'Стримы 4.1 и 4.2 остаются в практической разработке; LabVIEW назначен дополнительно по описаниям автора.']:summary.add_paragraph(t,'List Bullet')
summary.add_heading('5. Полное распределение по видео',1)
summary.add_paragraph('Ниже приведены исходные названия и кликабельные ссылки. Короткая надпись на новой обложке не означает автоматического переименования видео. Полный путь файла и дополнительные плейлисты — в Excel.')
for k,g in groups.items():
 summary.add_heading(f"{g['name']} — {len(g['nums'])} видео",1);summary.add_paragraph(g['description']);summary.add_paragraph('Рекомендуемый порядок: '+g['sorting'])
 for pos,n in enumerate(orders[k],1):
  a=rows[n-1];link(summary,f"{pos:02}. [№{n:02}] {a['title']} — {a['id']}",a['url'])
summary.add_heading('6. Дополнительные подборки',1)
extra_map=collections.defaultdict(list)
for a in rows:
 for p in a['extra'].split('; '):
  if p:extra_map[p].append(a)
for p,rs in extra_map.items():
 summary.add_heading(p,2);summary.add_paragraph('Добавить дополнительно к основному назначению:')
 for a in rs:link(summary,f"[№{a['number']:02}] {a['title']}",a['url'])
summary.add_heading('7. Перед загрузкой и созданием плейлистов',1)
for t in ['Выбрать основную версию стрима №3: l3xewCeguhM или Szt2Qw5Gb7o. Обложки сделаны для обоих; не обязательно ставить оба в одну подборку.', 'MONITOR BALL имеет русскую и английскую версии. Сохранить обе, поставить рядом и оставить метку ENGLISH VERSION у английской.', 'LabVIEW 2018, MOKO SE версии 2 и инструкция Windows 10 — архивные материалы. Не выдавать их за актуальные версии ПО.', 'Shorts: подготовленные 1080×1920 — вертикальные кадры/макеты. Их нельзя считать обычными загружаемыми пользовательскими миниатюрами. Проверить доступный выбор кадра в текущем интерфейсе YouTube; для применения нового изображения может понадобиться включение его в сам ролик.', 'На части старых фото чёткость ограничена исходной обложкой. Особенно проверить станок, старые стримы, телеметрию и архивные проекты; улучшать их по оригинальным кадрам, не заменяя аппаратуру похожей генерацией.', 'Сначала посмотреть новые 76 файлов в галерее, затем менять обложки партиями. Сохранённые 18 не требуют повторного дизайна, если нет новых замечаний.', 'Старые плейлисты не удалять до проверки переноса всех роликов. «Новинки» и «СМИ о нас» можно оставить дополнительными подборками.']:summary.add_paragraph(t,'List Bullet')
summary.add_heading('8. Позиции вне публичного списка',1)
table(summary,['ID','Что известно','Действие'],[['7ayoYanUV-w','Название недоступно; старый плейлист драйверов','Проверить доступность'],['Y9rNN698aM8','USB with LabVIEW; только в старом плейлисте','Уточнить автора'],['_FvfYFL_1vM','Tibo 2021 — возможность заглянуть в будущее','Уточнить внешнюю публикацию']])
summary.add_heading('9. Источники',1)
for u in ['https://www.youtube.com/@mokoserver/videos','https://www.youtube.com/@mokoserver/shorts','https://www.youtube.com/@mokoserver/playlists','https://moko.by/']:summary.add_paragraph(u)
summary.add_paragraph('Также использованы предоставленный владельцем брендбук, SVG и фотографии, история согласования и описания отдельных видео. Дата сборки списка: 09.09.2026.')
summary.save(D/'MOKO_Итоговая_карта_плейлистов.docx')
# Final rules retain the detailed agreed standard and add actual full-run decisions.
rules=Document(R/'docs/MOKO_Правила_обложек_v1.1.docx')
for p in rules.paragraphs:
 if 'версия 1.1' in p.text:p.text=p.text.replace('версия 1.1','версия 2.0 — полный выпуск')
 if 'Сами новые обложки ещё подлежат проверке владельцем.' in p.text:p.text=p.text.replace('Сами новые обложки ещё подлежат проверке владельцем.','Эта партия вошла в первые 18 согласованных образцов.')
rules.add_heading('14. Применение стандарта в полном выпуске',1)
for t in ['Основные цвета сохранены. Дополнительно применены бирюзовый #55D6C2 для инструкций MOKO SE и фиолетовый #B69AFF для медиа. Эти расширения представлены в новом комплекте для окончательного просмотра.', 'Интерфейсы и фотографии помещаются в рамку с размером, подстроенным под пропорции источника. Не растягивать снимок и не оставлять огромные пустые поля внутри рамки. Старую коралловую окантовку обрезать.', 'Для реальных устройств допустима обычная коррекция экспозиции тёмного фото, но не генеративное дорисовывание модели. В частности, фотография станка осветлена, а не заменена иллюстрацией другого станка.', 'Когда фото является документальным кадром с человеком и оборудованием, можно сохранить его целиком в рамке. Это отличается от неудачного узкого прямоугольного вырезания одного лица. Не обрезать человеку глаза или макушку ради выделения прибора.', 'Каждый новый файл сопоставлен с ID и брифом: тема, подзаголовок, формат, источник визуала. Полный комплект не означает автоматическое утверждение всех новых вариантов: 18 были согласованы раньше, 76 представлены на финальную проверку.', 'Вертикальные макеты Shorts выпускаются отдельно. В них можно адаптировать верхний бренд-блок под узкий формат, сохранив знак и название MOKO / SYSTEMS в композиции.']:rules.add_paragraph(t)
rules.save(D/'MOKO_Правила_обработки_обложек_v2.docx')
(F/'docs/README.txt').write_text('Полный выпуск: 92 JPG 1280×720 + 2 вертикальных макета 1080×1920.\n18 прежних файлов сохранены, 76 новых — для финальной проверки.\nExcel: все видео, плейлисты, дополнительные назначения, проверки и контроль SHA-256.\nWord: итоговая карта с каждым видео и подробные правила обработки.\nНи один файл не загружался на YouTube.\n')
print('Final docs created. Categories:',{k:len(g['nums']) for k,g in groups.items()});print('Files verified:',len(rows))
