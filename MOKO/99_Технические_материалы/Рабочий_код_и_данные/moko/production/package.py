from pathlib import Path
import json,base64,html,zipfile
from PIL import Image,ImageDraw,ImageFont
from openpyxl import load_workbook
from docx import Document
R=Path('/home/user/moko');P=R/'production';rows=json.load(open(R/'catalog/registry.json'))
order=['b5zjy0PHU0s','EsLUSn9RDL4','tXsjlxHMt6Y','nnktD99huyE','feEv04DYWbA','ex8ibsNSCiA']+[a['id'] for a in json.load(open(R/'batch02/new_items.json'))]
lookup={a['id']:a for a in rows}
def f(size,w=500):
 q=ImageFont.truetype(str(R/'brand/Inter.ttf'),size);q.set_variation_by_axes([14,w]);return q
sheet=Image.new('RGB',(1968,1910),'#080d15');d=ImageDraw.Draw(sheet);d.text((36,26),'MOKO / ПЕРВЫЕ 12 ОБЛОЖЕК',font=f(32,650),fill='white');d.text((36,79),'01–06 — СОГЛАСОВАННЫЕ / 07–12 — НОВАЯ ПАРТИЯ',font=f(19),fill='#a3aebe')
cards=[]
for i,vid in enumerate(order):
 a=lookup[vid];p=P/'thumbnails'/(vid+'.jpg');im=Image.open(p);assert im.size==(1280,720) and p.stat().st_size<2_000_000
 x=36+(i%3)*644;y=135+(i//3)*433;sheet.paste(im.resize((608,342),Image.Resampling.LANCZOS),(x,y));d.text((x,y+353),f'{i+1:02} / '+a['headline'],font=f(22,600),fill='white');d.text((x,y+390),a['language'] if a['key'] in ['PY','LV'] else a['visual'],font=f(16),fill=a['color'])
 uri='data:image/jpeg;base64,'+base64.b64encode(p.read_bytes()).decode()
 cards.append(f'<article><button class="pic" onclick="show({i})"><img src="{uri}" alt="{html.escape(a["headline"])}"></button><h2>{i+1:02} / {html.escape(a["headline"])}</h2><p class="tag" style="color:{a["color"]}">{html.escape(a["visual"])}</p><p>{"Согласованный образец" if i<6 else "Новая партия — проверить"}</p><button onclick="show({i})">Открыть крупно</button></article>')
sheet.save(P/'all_12_preview.jpg',quality=95)
page='''<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>MOKO · Первые 12 обложек</title><style>*{box-sizing:border-box}body{margin:0;background:#080d15;color:#f5f7fb;font:15px/1.5 system-ui,sans-serif}header,main,footer{max-width:1500px;margin:auto;padding:26px}h1{font-size:clamp(28px,4vw,44px);line-height:1.1}p{color:#a3aebe}header p{max-width:850px}.k{color:#ed6262;letter-spacing:2px;font-size:12px}main{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:28px 22px}button{font:inherit;color:white;background:#172233;border:1px solid #334156;border-radius:9px;padding:8px 12px;cursor:pointer}.pic{display:block;width:100%;padding:0;border:0;overflow:hidden}img{display:block;width:100%}h2{font-size:19px;margin:12px 0 4px}.tag{margin:0;font-size:13px}article p{font-size:13px;margin-top:5px}dialog{background:#0b1220;border:1px solid #334156;color:white;border-radius:12px;width:min(1320px,96vw);max-height:95vh;padding:15px}dialog::backdrop{background:#000d}.dh{display:flex;gap:20px;justify-content:space-between;align-items:center;margin-bottom:14px}.dh h2{margin:0}a{color:#a7c9ff}.foot{font-size:12px}@media(max-width:1050px){main{grid-template-columns:1fr 1fr}}@media(max-width:640px){main{grid-template-columns:1fr}header,main,footer{padding:18px}}</style><header><div class="k">MOKO / SYSTEMS · ВЫПУСК 01</div><h1>Первые 12 обложек</h1><p>Шесть согласованных образцов и шесть новых. Python — зелёный, LabVIEW — оранжевый; язык подписан на учебных обложках. Полный реестр содержит 94 ролика, но здесь показаны только подготовленные файлы.</p></header><main>'''+''.join(cards)+'''</main><footer><p class="foot">Портреты на LED-фаре, плагине и выставке стилизованы генеративно. В плагине и State Machine используются реальные интерфейс/схема. Остальные технические изображения — концептуальные иллюстрации, не точные скриншоты. Файл работает без интернета.</p></footer><dialog id="dlg"><div class="dh"><h2 id="title"></h2><button onclick="document.getElementById('dlg').close()">×</button></div><img id="img"><p><a id="dl">Скачать JPG</a></p></dialog><script>const titles='''+json.dumps([lookup[v]['headline'] for v in order],ensure_ascii=False)+''';const ids='''+json.dumps(order)+''';function show(i){const src=document.querySelectorAll('.pic img')[i].src;document.getElementById('title').textContent=titles[i];document.getElementById('img').src=src;document.getElementById('dl').href=src;document.getElementById('dl').download=ids[i]+'.jpg';document.getElementById('dlg').showModal()}</script></html>'''
(P/'preview_12.html').write_text(page)
readme='''MOKO Systems — выпуск 01, 09.09.2026

Подготовлено 12 горизонтальных обложек: 6 согласованных образцов, 6 новых.
Всего в реестре 94 доступных ролика: 92 обычных + 2 Shorts.
Осталось 80 обычных обложек и 2 вертикальных макета Shorts.

thumbnails/ — файлы названы по YouTube ID.
preview_12.html — просмотр 12 обложек на телефоне, без интернета.
all_12_preview.jpg — общий лист.
docs/ — правила, сводная карта плейлистов, Excel и CSV.

Перед загрузкой проверить новые 6 вариантов. Сами файлы на YouTube не загружены.
Сводная карта — распределение всех 94 роликов по метаданным, не отчёт о готовности всех обложек.
'''
(P/'README.txt').write_text(readme)
with zipfile.ZipFile(P/'MOKO_Выпуск_01_12_обложек_и_документы.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in sorted((P/'thumbnails').glob('*.jpg')):z.write(p,'thumbnails/'+p.name)
 for p in sorted((R/'docs').glob('*')):
  if p.suffix in ['.docx','.xlsx','.csv','.md']:z.write(p,'docs/'+p.name)
 for name in ['preview_12.html','all_12_preview.jpg','README.txt']:z.write(P/name,name)
 z.write(R/'brand/logo.svg','brand/logo.svg')
# Structural validation of deliverables.
wb=load_workbook(R/'docs/MOKO_Реестр_94_видео_и_плейлисты.xlsx');assert wb['Видео'].max_row==95;assert wb['Готовые файлы'].max_row==13
for p in (R/'docs').glob('*.docx'):
 q=Document(p);print(p.name,'paragraphs',len(q.paragraphs),'tables',len(q.tables))
print('12 image sizes verified, workbook 94 rows verified. ZIP bytes', (P/'MOKO_Выпуск_01_12_обложек_и_документы.zip').stat().st_size)
